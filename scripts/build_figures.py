"""Fill the briefing template with figures from the dashboard manifest, the OBR's
detailed forecast table 1.1 and the Treasury's Tables 6.1 and 6.2.

Every PolicyEngine figure comes from a manifest artifact whose files match the
digests in its receipt.json; every official figure comes from the OBR workbook
(checked against its sha256) or the Table 6.1 workbook. figures.json records the
source of each figure printed.

Official comparisons use the OBR line or lines that match what we model; Table
6.1's rounded total is context only (Max, 28 Oct 2026). <lines> joins OBR
category names with "+", e.g. Income tax+NICs. <n> is the OBR measure number,
the same as its Table 6.1 row.

Placeholders in the template:
  {{costing:<measure_id>:<n>:<lines>[:<label>]}}  table: OBR line(s), Table 6.1 total, the dashboard's estimate
                                  (its row labelled <label>, "PolicyEngine" if none)
  {{official:<n>:<lines>}}        table: OBR line(s) and Table 6.1 total
  {{official:<n>}}                table: OBR measure total only
  {{obr:<n>:<year>}}              one OBR measure total, signed, £m
  {{obrabs:<n>:<year>}}           one OBR measure total as £ without sign, £m
  {{line:<n>:<lines>:<year>}}     one OBR line value, signed, £m
  {{lineabs:<n>:<lines>:<year>}}  one OBR line value as £ without sign, £m
  {{gap:<measure_id>:<n>:<lines>:<year>}}     |ours − OBR line| as £ without sign, £m, from the printed whole £m
  {{gappct:<measure_id>:<n>:<lines>:<year>}}  |ours ÷ OBR line − 1| as a whole per cent, from the printed whole £m
  {{t61:<row>:<year>}}            one Table 6.1 value, signed, £m
  {{t61abs:<row>:<year>}}         one Table 6.1 value as £ without sign, £m
  {{t62sum:<row>}}                a Table 6.2 row summed over the six years, as £ without sign, £m
  {{pe:<measure_id>:<year>}}      one dashboard value, signed, £m, rounded
  {{meta:<key>}}                  manifest revision, updated_at, source URL, source label
  {{signoff:<combined key>}}      nothing once the combined run carries both signatures; until then a
                                  sentence, with a leading space, saying the combined figure may change
  {{status_note}}                 the opening note: provisional or approved, test or production
  {{keypoints:<post slug>[:<n>]}} a live explainer's key points as published, and n paragraphs after them
  {{post_sha}}                    the policyengine-app-v2 commit the saved explainers come from
"""

import argparse
import csv
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))  # python -I leaves out the script's folder
from explainers import POSTS, key_points  # noqa: E402

YEARS = [2026, 2027, 2028, 2029, 2030, 2031]
MINUS = "−"
DASHBOARD = "[policyengine.org/uk/autumn-budget-2026](https://www.policyengine.org/uk/autumn-budget-2026)"
# policy.xlsx, the OBR's detailed forecast table 1.1 for October 2026, as checked into dashboard #300 at 555e7f74
OBR_SHA256 = "cdcc2a06b100241bf3010bd159a4e5d43daef2348a3850797b1e88a308fd45fd"


def signed(attestation):
    """Both signatures present, neither a preview placeholder."""
    att = attestation or {}
    return all(isinstance(v, str) and v and "PREVIEW ONLY" not in v for v in (att.get("checker"), att.get("second")))


def year_label(y):
    return f"{y}-{str(y + 1)[-2:]}"


def fmt(v):
    """Whole £m, signed: −575, +35, 0."""
    n = round(v)
    if n == 0:
        return "0"
    s = f"{abs(n):,}"
    return f"{MINUS}{s}" if n < 0 else f"+{s}"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_table(xlsx, sheet, first_value_col):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[sheet]
    rows = {}
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[1], int) and isinstance(r[2], str):
            vals = r[first_value_col:first_value_col + len(YEARS)]
            if all(isinstance(v, (int, float)) for v in vals):
                rows[r[1]] = {"title": r[2], "values": dict(zip(YEARS, vals))}
    return rows


def load_obr(xlsx):
    """Measure number -> {"title", "lines": {category: {year: £m}}, "total": {year: £m}}."""
    if sha(xlsx) != OBR_SHA256:
        raise SystemExit(f"{xlsx}: sha256 is not the checked OBR workbook's")
    rows = list(openpyxl.load_workbook(xlsx, data_only=True)["1.1"].iter_rows(values_only=True))
    out = {}
    for i, r in enumerate(rows):
        if not (isinstance(r[2], int) and isinstance(r[3], str)) or r[2] in out:
            continue
        lines = {}
        for row in rows[i:]:
            vals = dict(zip(YEARS, row[8:14]))
            if row[6] == "Total":
                out[r[2]] = {"title": r[3].strip(), "lines": lines, "total": vals}
                break
            lines[row[7]] = {y: lines.get(row[7], {}).get(y, 0) + vals[y] for y in YEARS}
    return out


def load_dashboard(manifest_path, art_dir):
    m = json.loads(Path(manifest_path).read_text())
    out = {}
    entries = {**m["measures"], **m["combined"]}
    for key, entry in entries.items():
        folder = Path(art_dir) / key
        if not folder.exists():
            continue
        receipt = json.loads((folder / "receipt.json").read_text())["dashboard"]
        for name, digest in receipt["files"].items():
            got = sha(folder / name)
            if got != digest:
                raise SystemExit(f"{key}/{name}: digest {got} != receipt {digest}")
        if entry.get("outputs_sha256") and receipt["digests"]["outputs_sha256"] != entry["outputs_sha256"]:
            raise SystemExit(f"{key}: receipt outputs digest differs from the manifest")
        rows = list(csv.DictReader(open(folder / "budgetary_impact.csv")))
        out[key] = {
            "status": entry.get("status"),
            "artifact": entry["artifact"],
            "values_bn": {int(r["year"]): float(r["value"]) for r in rows},
        }
    return m, out


def costing_table(official_rows, pe=None, pe_label="PolicyEngine"):
    head = "| £ million | " + " | ".join(year_label(y) for y in YEARS) + " |"
    sep = "|:--------------------------|" + "|".join(["------:"] * len(YEARS)) + "|"
    lines = [head, sep]
    for label, vals in official_rows:
        lines.append(f"| {label} | " + " | ".join(fmt(vals[y]) for y in YEARS) + " |")
    if pe is not None:
        lines.append(f"| {pe_label} | " + " | ".join(fmt(pe[y]) for y in YEARS) + " |")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--manifest-url", required=True)
    ap.add_argument("--artifacts", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--obr", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--figures", required=True)
    ap.add_argument("--source", choices=["test", "production"], default="test")
    a = ap.parse_args()

    # Table 6.1 has two head columns before the years; Table 6.2 has one.
    t61 = load_table(a.xlsx, "Table 6.1", 5)
    t62 = load_table(a.xlsx, "Table 6.2", 4)
    obr = load_obr(a.obr)
    manifest, dash = load_dashboard(a.manifest, a.artifacts)
    figures = []

    def obr_line(n, lines):
        """The sum of the named lines of OBR measure n; every name must exist."""
        names = lines.split("+")
        missing = [x for x in names if x not in obr[n]["lines"]]
        if missing:
            raise SystemExit(f"OBR measure {n} has no line {missing}; lines: {list(obr[n]['lines'])}")
        return {y: sum(obr[n]["lines"][x][y] for x in names) for y in YEARS}

    def line_label(lines):
        return " and ".join(lines.split("+"))

    def others_note(n, lines):
        """Name the measure's other non-zero OBR lines, which Table 6.1's total also counts."""
        names = lines.split("+")
        rest = [c for c, v in obr[n]["lines"].items() if c not in names and any(round(x, 6) for x in v.values())]
        if not rest:
            return ""
        listed = rest[0] if len(rest) == 1 else ", ".join(rest[:-1]) + " and " + rest[-1]
        return f"\n\nTable 6.1's total also includes the measure's other {'line' if len(rest) == 1 else 'lines'} in the OBR's table: {listed}."

    def pe_m(key):
        if key not in dash:
            raise SystemExit(f"no verified artifact for {key}")
        if dash[key]["status"] == "withdrawn":
            raise SystemExit(f"{key} is withdrawn from the release; no PolicyEngine figure may print")
        return {y: dash[key]["values_bn"][y] * 1000 for y in YEARS}

    def note(kind, key, year, value, printed):
        if kind == "pe":
            src = {"file": f"{dash[key]['artifact']}/budgetary_impact.csv", "row": f"year={year}", "column": "value (£bn)"}
        elif kind == "obr":
            src = {"file": f"policy-555e7f74.xlsx (sha256 {OBR_SHA256[:8]}), sheet 1.1", "row": f"measure {key}", "column": year_label(year)}
        elif kind == "gap":
            src = {"file": "computed", "row": key, "column": year_label(year)}
        else:
            src = {"file": "Table_6.1_and_6.2_Budget_2026_Policy_Decisions.xlsx", "row": f"Table {'6.2' if kind == 't62' else '6.1'} row {key}", "column": year_label(year) if year else "sum 2026-27 to 2031-32"}
        figures.append({"kind": kind, "key": str(key), "year": year_label(year) if year else None, "value": value, "printed": printed, **src})

    def note_rows(kind, key, vals):
        for y in YEARS:
            note(kind, key, y, vals[y], fmt(vals[y]))

    def sub(match):
        parts = match.group(1).split(":")
        kind = parts[0]
        if kind == "costing":
            key, n, lines = parts[1], int(parts[2]), parts[3]
            pe_label = parts[4] if len(parts) > 4 else "PolicyEngine"
            pe, line = pe_m(key), obr_line(n, lines)
            note_rows("obr", f"{n}, {lines}", line)
            note_rows("t61", n, t61[n]["values"])
            note_rows("pe", key, pe)
            return costing_table([(f"Official: {line_label(lines)}", line), ("Table 6.1 total", t61[n]["values"])], pe, pe_label) + others_note(n, lines)
        if kind == "official" and len(parts) == 3:
            n, lines = int(parts[1]), parts[2]
            line = obr_line(n, lines)
            note_rows("obr", f"{n}, {lines}", line)
            note_rows("t61", n, t61[n]["values"])
            return costing_table([(f"Official: {line_label(lines)}", line), (f"Table 6.1 total (row {n})", t61[n]["values"])]) + others_note(n, lines)
        if kind == "official":
            n = int(parts[1])
            note_rows("obr", f"{n}, Total", obr[n]["total"])
            return costing_table([(f"Official (measure {n})", obr[n]["total"])])
        if kind in ("obr", "obrabs"):
            n, y = int(parts[1]), int(parts[2])
            v = obr[n]["total"][y]
            printed = fmt(v) if kind == "obr" else f"£{abs(round(v)):,}"
            note("obr", f"{n}, Total", y, v, printed)
            return printed
        if kind in ("line", "lineabs"):
            n, lines, y = int(parts[1]), parts[2], int(parts[3])
            v = obr_line(n, lines)[y]
            printed = fmt(v) if kind == "line" else f"£{abs(round(v)):,}"
            note("obr", f"{n}, {lines}", y, v, printed)
            return printed
        if kind in ("gap", "gappct"):
            key, n, lines, y = parts[1], int(parts[2]), parts[3], int(parts[4])
            # From the whole £m printed in the card's table, so a reader can check it by subtraction
            ours, theirs = round(pe_m(key)[y]), round(obr_line(n, lines)[y])
            if kind == "gap":
                v = abs(ours - theirs)
                printed = f"£{round(v):,}"
            else:
                v = abs(ours / theirs - 1) * 100
                printed = f"{round(v)}%"
            note("gap", f"{kind} {key} vs OBR {n}, {lines}", y, v, printed)
            return printed
        if kind == "t61":
            row, y = int(parts[1]), int(parts[2])
            v = t61[row]["values"][y]
            note("t61", row, y, v, fmt(v))
            return fmt(v)
        if kind == "t61abs":
            row, y = int(parts[1]), int(parts[2])
            v = t61[row]["values"][y]
            printed = f"£{abs(round(v)):,}"
            note("t61", row, y, v, printed)
            return printed
        if kind == "t62sum":
            row = int(parts[1])
            v = sum(t62[row]["values"][y] for y in YEARS)
            printed = f"£{abs(round(v)):,}"
            note("t62", row, None, v, printed)
            return printed
        if kind == "pe":
            key, y = parts[1], int(parts[2])
            v = pe_m(key)[y]
            note("pe", key, y, v, fmt(v))
            return fmt(v)
        if kind == "peabs":
            key, y = parts[1], int(parts[2])
            v = pe_m(key)[y]
            printed = f"£{abs(round(v)):,}"
            note("pe", key, y, v, printed)
            return printed
        if kind == "meta":
            return {
                "revision": str(manifest["revision"]),
                "updated_at": manifest["updated_at"][:16].replace("T", " ") + " UTC",
                "url": a.manifest_url,
                "source_label": "its test publication manifest" if a.source == "test" else "its publication manifest",
            }[parts[1]]
        if kind == "signoff":
            entry = manifest["combined"][parts[1]]
            if signed(entry.get("attestation")):
                return ""
            return " The combined figure is provisional and may change."
        if kind == "post_sha":
            return (POSTS / "APP_V2_MAIN_SHA").read_text().strip()
        if kind == "keypoints":
            return key_points(parts[1], int(parts[2]) if len(parts) > 2 else 0)
        if kind == "status_note":
            shown = [k for k, v in manifest["measures"].items() if v.get("status") in ("approved", "provisional")]
            provisional = any(manifest["measures"][k]["status"] == "provisional" for k in shown)
            rev = f"revision {manifest['revision']} ({manifest['updated_at'][:16].replace('T', ' ')} UTC)"
            if a.source == "test":
                lead = f"PolicyEngine figures are first estimates from the Budget dashboard's test publication, {rev}."
                tail = f" They may change when we publish the dashboard at {DASHBOARD}."
            else:
                lead = f"Our figures come from our Budget dashboard at {DASHBOARD}, as it stood at {rev}."
                tail = " The dashboard still marks some as provisional, and they may change." if provisional else ""
            flag = "[Provisional]{.flag} " if provisional else ""
            return flag + lead + tail
        raise SystemExit(f"unknown placeholder {match.group(0)}")

    text = Path(a.template).read_text()
    filled = re.sub(r"\{\{([^}]+)\}\}", sub, text)
    Path(a.out).write_text(filled)
    Path(a.figures).write_text(json.dumps(
        {"manifest_url": a.manifest_url, "revision": manifest["revision"], "updated_at": manifest["updated_at"], "figures": figures},
        indent=1, ensure_ascii=False))
    print(f"revision {manifest['revision']}: {len(figures)} figures written to {a.figures}")


if __name__ == "__main__":
    main()
